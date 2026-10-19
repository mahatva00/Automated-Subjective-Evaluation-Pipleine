"""Parse an assessment policy without inventing unstated scoring rules.

Supported inputs: JSON, TXT, Markdown, PDF and DOCX. The output preserves the
source text/hash and only enables numeric scoring when credit factors are
explicitly supplied as numeric values.
"""
from __future__ import annotations
import hashlib
import json
import re
from pathlib import Path
from typing import Any

VALID_STATUSES = ("complete", "partial", "absent")
POLICY_SUFFIXES = {".json", ".txt", ".md", ".markdown", ".pdf", ".docx"}
COMPONENT_NAMES = ("code", "explanation", "output")


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _pdf_text(path: Path, input_root: Path | None) -> tuple[str, list[str]]:
    from src.ingestion.pdf_parser import parse_pdf
    root = (input_root or path.parent).resolve()
    result = parse_pdf(path.resolve(), root)
    segments = result.get("segments", []) or []
    text = "\n\n".join(str(s.get("content") or "") for s in segments if isinstance(s, dict))
    diagnostics = list(result.get("diagnostics", []) or [])
    return text, [str(x) for x in diagnostics]


def _docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Reading DOCX policies requires python-docx. Install it with `pip install python-docx`.") from exc
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            vals = [cell.text.strip() for cell in row.cells]
            if any(vals):
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def _extract_explicit_credit_map(payload: Any, text: str) -> dict[str, float] | None:
    candidate = None
    if isinstance(payload, dict):
        for key in ("credit_map", "status_credits", "status_credit", "credit_factors"):
            if isinstance(payload.get(key), dict):
                candidate = payload[key]
                break
        if candidate is None:
            grading = payload.get("grading") or payload.get("scoring")
            if isinstance(grading, dict):
                for key in ("credit_map", "status_credits", "credit_factors"):
                    if isinstance(grading.get(key), dict):
                        candidate = grading[key]
                        break
    if candidate is not None:
        result: dict[str, float] = {}
        for status in VALID_STATUSES:
            value = candidate.get(status)
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                return None
            number = float(value)
            if not 0.0 <= number <= 1.0:
                return None
            result[status] = number
        return result

    # Text policies are only normalized automatically when all factors are stated numerically.
    result = {}
    for status in VALID_STATUSES:
        pattern = re.compile(
            rf"\b{status}\b\s*(?:answer|response|work)?\s*(?:=|:|->|receives?|earns?|gets?)\s*"
            rf"(\d+(?:\.\d+)?)\s*(%|percent)?(?=\s|[,;.]|$)", re.I
        )
        match = pattern.search(text)
        if not match:
            return None
        value = float(match.group(1))
        if match.group(2) or value > 1.0:
            value /= 100.0
        if not 0.0 <= value <= 1.0:
            return None
        result[status] = value
    return result


def _strictness(payload: Any, text: str) -> str:
    if isinstance(payload, dict):
        raw = payload.get("strictness") or payload.get("grading_style")
        if isinstance(raw, str) and raw.strip().lower() in {"strict", "lenient", "balanced"}:
            return raw.strip().lower()
    found = {x for x in ("strict", "lenient", "balanced") if re.search(rf"\b{x}\b", text, re.I)}
    return next(iter(found)) if len(found) == 1 else "unspecified"


def _normalise_component_weights(values: dict[str, Any]) -> dict[str, float] | None:
    result: dict[str, float] = {}
    for component in COMPONENT_NAMES:
        value = values.get(component)
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            return None
        number = float(value)
        if number > 1.0:
            number /= 100.0
        if not 0.0 <= number <= 1.0:
            return None
        result[component] = number
    if abs(sum(result.values()) - 1.0) > 0.001:
        return None
    return result


def _extract_subquestion_marks(section: str, question_id: str) -> dict[str, float]:
    row_pattern = re.compile(
        rf"^\s*{re.escape(question_id)}\(([a-z])\)\s+(?:[A-Z]\d+\s+)?"
        r"(\d+(?:\.\d+)?)\s+",
        re.I | re.M,
    )
    return {
        f"{question_id}({match.group(1).lower()})": float(match.group(2))
        for match in row_pattern.finditer(section)
    }


def _extract_component_policy(payload: Any, text: str) -> tuple[dict[str, dict[str, Any]], list[str]]:
    """Extract explicit question/component allocations without guessing missing marks."""
    questions: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []

    if isinstance(payload, dict):
        raw_questions = payload.get("questions")
        if isinstance(raw_questions, dict):
            question_items = raw_questions.items()
        elif isinstance(raw_questions, list):
            question_items = (
                (str(item.get("question_id") or item.get("id")), item)
                for item in raw_questions
                if isinstance(item, dict)
            )
        else:
            question_items = ()
        for question_id, item in question_items:
            if not isinstance(item, dict):
                continue
            weights = item.get("components") or item.get("component_weights")
            normalised = _normalise_component_weights(weights) if isinstance(weights, dict) else None
            marks = item.get("marks")
            if normalised is not None and isinstance(marks, (int, float)) and not isinstance(marks, bool) and marks >= 0:
                questions[str(question_id)] = {
                    "marks": float(marks),
                    "component_weights": normalised,
                    "subquestion_marks": {},
                }
        if questions:
            return questions, warnings

    heading_pattern = re.compile(
        r"^\s*Q\s*(\d+)\b.*?\[\s*(\d+(?:\.\d+)?)\s*marks?\s*\]\s*$",
        re.I | re.M,
    )
    split_pattern = re.compile(
        r"^\s*Split\s*:\s*Code\s*(\d+(?:\.\d+)?)\s*%\s*,\s*"
        r"Explanation\s*(\d+(?:\.\d+)?)\s*%\s*,\s*Output\s*(\d+(?:\.\d+)?)\s*%\s*$",
        re.I | re.M,
    )
    headings = list(heading_pattern.finditer(text))
    for index, heading in enumerate(headings):
        question_id = heading.group(1)
        section_end = headings[index + 1].start() if index + 1 < len(headings) else len(text)
        section = text[heading.end():section_end]
        split = split_pattern.search(section)
        if split is None:
            warnings.append(f"Question Q{question_id} has no complete Code/Explanation/Output split.")
            continue
        weights = _normalise_component_weights({
            "code": float(split.group(1)),
            "explanation": float(split.group(2)),
            "output": float(split.group(3)),
        })
        if weights is None:
            warnings.append(f"Question Q{question_id} has component weights that do not sum to 100%.")
            continue
        questions[question_id] = {
            "marks": float(heading.group(2)),
            "component_weights": weights,
            "subquestion_marks": _extract_subquestion_marks(section, question_id),
        }
    return questions, warnings


def parse_assessment_policy(path: str | Path | None, input_root: str | Path | None = None,
                            parsed_pdf: dict[str, Any] | None = None) -> dict[str, Any]:
    """Return a normalized policy; unresolved scoring remains explicitly unresolved."""
    if path is None:
        return {
            "schema_version": "1.0", "artifact_type": "normalized_assessment_policy",
            "source_filename": None, "source_sha256": None, "source_format": None,
            "strictness": "unspecified", "credit_map": None,
            "numeric_scoring_enabled": False, "mark_allocation": {"mode": "per_source_node_equal_split"},
            "question_component_policy": {}, "component_scoring_enabled": False,
            "teacher_approval_required": True, "raw_text": "", "warnings": ["No assessment policy file was identified; numeric grading is disabled."],
        }
    source = Path(path).resolve()
    suffix = source.suffix.lower()
    if suffix not in POLICY_SUFFIXES:
        raise ValueError(f"Unsupported assessment policy format {suffix!r}: {source.name}")
    payload: Any = None
    diagnostics: list[str] = []
    if suffix == ".json":
        try:
            payload = json.loads(source.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ValueError(f"Assessment policy JSON is invalid: {exc}") from exc
        if not isinstance(payload, dict):
            raise ValueError("Assessment policy JSON must contain an object at its root.")
        raw_text = "\n".join(str(payload.get(k)) for k in ("policy_text", "instructions", "description") if payload.get(k) is not None)
        if not raw_text:
            raw_text = json.dumps(payload, ensure_ascii=False, indent=2)
    elif suffix == ".pdf":
        if parsed_pdf is not None:
            raw_text = "\n\n".join(str(s.get("content") or "") for s in parsed_pdf.get("segments", []) if isinstance(s, dict))
            diagnostics.extend(str(x) for x in parsed_pdf.get("diagnostics", []) or [])
        else:
            raw_text, diagnostics = _pdf_text(source, Path(input_root).resolve() if input_root else None)
    elif suffix == ".docx":
        raw_text = _docx_text(source)
    else:
        raw_text = source.read_text(encoding="utf-8-sig")

    credit_map = _extract_explicit_credit_map(payload, raw_text)
    question_component_policy, component_warnings = _extract_component_policy(payload, raw_text)
    mark_allocation: dict[str, Any] = {"mode": "per_source_node_equal_split"}
    teacher_approval_required = True
    if isinstance(payload, dict):
        nested = payload.get("grading") or payload.get("scoring")
        nested = nested if isinstance(nested, dict) else {}
        raw_alloc = payload.get("mark_allocation", nested.get("mark_allocation"))
        if isinstance(raw_alloc, dict) and raw_alloc.get("mode") in {"per_source_node_equal_split", "equal_within_question", "explicit_only"}:
            mark_allocation = {"mode": raw_alloc["mode"]}
        elif isinstance(raw_alloc, str) and raw_alloc in {"per_source_node_equal_split", "equal_within_question", "explicit_only"}:
            mark_allocation = {"mode": raw_alloc}
        val = payload.get("teacher_approval_required", nested.get("teacher_approval_required", True))
        teacher_approval_required = val if isinstance(val, bool) else True

    warnings = list(diagnostics) + component_warnings
    if credit_map is None:
        warnings.append("No complete numeric credit map for complete/partial/absent was found; only fulfillment summaries will be reported, not numeric marks.")
    return {
        "schema_version": "1.0", "artifact_type": "normalized_assessment_policy",
        "source_filename": source.name, "source_path": str(source), "source_sha256": _sha256(source),
        "source_format": suffix.lstrip("."), "strictness": _strictness(payload, raw_text),
        "credit_map": credit_map, "numeric_scoring_enabled": credit_map is not None,
        "question_component_policy": question_component_policy,
        "component_scoring_enabled": bool(question_component_policy),
        "mark_allocation": mark_allocation, "teacher_approval_required": teacher_approval_required,
        "raw_text": raw_text, "warnings": warnings,
    }
